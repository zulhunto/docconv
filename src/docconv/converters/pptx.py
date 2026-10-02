"""PowerPoint (.pptx) document converter."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import ClassVar

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from docconv.assets import AssetManager
from docconv.converters.base import BaseConverter
from docconv.models import Asset, ConvertContext, DocMetadata, MarkdownResult
from docconv.postprocess.tables import TableCell, render_structured_table
from docconv.registry import register

logger = logging.getLogger(__name__)


@register
class PptxConverter(BaseConverter):
    """Converts PowerPoint presentations (.pptx) to Markdown."""

    extensions: ClassVar[tuple[str, ...]] = (".pptx",)

    def convert(self, path: Path, ctx: ConvertContext) -> MarkdownResult:
        warnings: list[str] = []
        assets: list[Asset] = []

        try:
            prs = Presentation(path)
        except Exception as exc:
            return MarkdownResult.from_error(f"Cannot open PPTX: {exc}")

        # Metadata
        cp = prs.core_properties
        title = cp.title or path.stem.replace("_", " ").title()
        metadata = DocMetadata(
            title=title,
            author=cp.author or "",
            created=cp.created.isoformat() if cp.created else "",
            page_count=len(prs.slides),
            source_format="PPTX",
        )

        asset_mgr = AssetManager(ctx.assets_dir, min_size_px=ctx.config.images.min_size_px)
        slides_md: list[str] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_parts: list[str] = []

            # Extract slide title
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()
                slide_parts.append(f"## Slide {slide_idx}: {slide_title}\n")
            else:
                slide_parts.append(f"## Slide {slide_idx}\n")

            # Extract shapes
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue  # already handled

                # Text box / shape with text
                if shape.has_text_frame:
                    paragraphs = []
                    for p in shape.text_frame.paragraphs:
                        p_txt = p.text.strip()
                        if not p_txt:
                            continue
                        # Indent level
                        prefix = "- " if p.level == 0 else f"{'  ' * p.level}- "
                        paragraphs.append(f"{prefix}{p_txt}")
                    if paragraphs:
                        slide_parts.append("\n".join(paragraphs))

                # Table
                elif shape.has_table:
                    tbl = shape.table
                    grid: list[list[TableCell]] = []
                    for r_idx, row in enumerate(tbl.rows):
                        grid_row = [
                            TableCell(text=c.text.strip(), is_header=(r_idx == 0))
                            for c in row.cells
                        ]
                        grid.append(grid_row)
                    table_md = render_structured_table(
                        grid,
                        fallback_html_on_merged=ctx.config.tables.fallback_html_on_merged_cells,
                    )
                    if table_md:
                        slide_parts.append(table_md)

                # Picture
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and ctx.extract_images:
                    try:
                        img_bytes = shape.image.blob
                        img_ext = shape.image.ext
                        asset = asset_mgr.save_image_bytes(
                            img_bytes,
                            hint_name=f"slide_{slide_idx}_img",
                            suffix=f".{img_ext}",
                        )
                        if asset:
                            assets.append(asset)
                            slide_parts.append(asset.markdown_ref)
                    except Exception as exc:
                        logger.debug("Failed to extract image on slide %d: %s", slide_idx, exc)

            # Speaker Notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(f"\n> **Speaker Notes:**\n> {notes.replace(chr(10), chr(10) + '> ')}")

            slides_md.append("\n\n".join(slide_parts))

        full_md = f"# {title}\n\n" + "\n\n---\n\n".join(slides_md) + "\n"
        return MarkdownResult(
            markdown=full_md,
            metadata=metadata,
            assets=assets,
            warnings=warnings,
        )
