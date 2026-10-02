"""Excel (.xlsx) and CSV/TSV document converter."""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import ClassVar

import openpyxl

from docconv.converters.base import BaseConverter
from docconv.models import ConvertContext, DocMetadata, MarkdownResult
from docconv.postprocess.tables import TableCell, render_pipe_table, render_structured_table
from docconv.registry import register

logger = logging.getLogger(__name__)


@register
class XlsxConverter(BaseConverter):
    """Converts Excel workbooks (.xlsx, .xlsm) and CSV/TSV files to Markdown."""

    extensions: ClassVar[tuple[str, ...]] = (".xlsx", ".xlsm", ".xltx", ".csv", ".tsv")

    def convert(self, path: Path, ctx: ConvertContext) -> MarkdownResult:
        ext = path.suffix.lower()
        if ext in (".csv", ".tsv"):
            return self._convert_csv(path, ctx)

        warnings: list[str] = []
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as exc:
            return MarkdownResult.from_error(f"Cannot open Excel file: {exc}")

        title = wb.properties.title or path.stem.replace("_", " ").title()
        metadata = DocMetadata(
            title=title,
            author=wb.properties.creator or "",
            created=wb.properties.created.isoformat() if wb.properties.created else "",
            source_format="Excel",
        )

        sheet_parts: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.max_row == 0 or ws.max_column == 0:
                continue

            # Build map of merged cell ranges
            merged_spans: dict[tuple[int, int], tuple[int, int]] = {}
            for merged_range in ws.merged_cells.ranges:
                min_col, min_row, max_col, max_row = merged_range.bounds
                colspan = max_col - min_col + 1
                rowspan = max_row - min_row + 1
                merged_spans[(min_row, min_col)] = (rowspan, colspan)

            grid: list[list[TableCell]] = []

            for r_idx in range(1, ws.max_row + 1):
                row_cells: list[TableCell] = []
                for c_idx in range(1, ws.max_column + 1):
                    # Check if covered by another merged cell
                    is_subsumed = False
                    for (m_r, m_c), (r_span, c_span) in merged_spans.items():
                        if m_r <= r_idx < m_r + r_span and m_c <= c_idx < m_c + c_span:
                            if (r_idx, c_idx) != (m_r, m_c):
                                is_subsumed = True
                            break

                    if is_subsumed:
                        continue

                    val = ws.cell(row=r_idx, column=c_idx).value
                    val_str = "" if val is None else str(val).strip()

                    rowspan, colspan = merged_spans.get((r_idx, c_idx), (1, 1))
                    row_cells.append(
                        TableCell(
                            text=val_str,
                            rowspan=rowspan,
                            colspan=colspan,
                            is_header=(r_idx == 1),
                        )
                    )

                # Skip completely blank rows
                if any(c.text for c in row_cells):
                    grid.append(row_cells)

            if grid:
                table_md = render_structured_table(
                    grid,
                    fallback_html_on_merged=ctx.config.tables.fallback_html_on_merged_cells,
                )
                sheet_parts.append(f"## Sheet: {sheet_name}\n\n{table_md}")

        full_md = f"# {title}\n\n" + "\n\n---\n\n".join(sheet_parts) + "\n"
        return MarkdownResult(markdown=full_md, metadata=metadata, warnings=warnings)

    def _convert_csv(self, path: Path, ctx: ConvertContext) -> MarkdownResult:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f, delimiter=delimiter)
                rows = list(reader)
        except Exception as exc:
            return MarkdownResult.from_error(f"Cannot read CSV file: {exc}")

        title = path.stem.replace("_", " ").title()
        metadata = DocMetadata(title=title, source_format="CSV")

        table_md = render_pipe_table(rows, has_header=True)
        full_md = f"# {title}\n\n{table_md}\n"
        return MarkdownResult(markdown=full_md, metadata=metadata)
